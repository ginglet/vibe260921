"""네이버 검색 결과에 나온 뉴스 기사의 본문을 수집하는 크롤러 (requests + BeautifulSoup4)

동작 순서
  1) 검색어(또는 검색 URL)로 네이버 뉴스 검색 결과를 요청해 결과 페이지에서 '네이버 뉴스' 기사 링크(n.news.naver.com)를 모은다.
  2) 기사 링크를 하나씩 열어 제목 / 언론사 / 기자 / 작성시각 / 본문을 추출한다.
  3) 결과를 JSON과 CSV(엑셀에서 한글이 깨지지 않는 utf-8-sig)로 저장한다. --xlsx 를 주면 엑셀 파일도 만든다.

사용법
  pip install requests beautifulsoup4 openpyxl       # openpyxl 은 엑셀 저장(--xlsx)을 쓸 때만 필요
  python naver_news_crawler.py 반도체                # 검색어만 입력해도 됩니다 (기사 10개)
  python naver_news_crawler.py "전기차 보조금"        # 띄어쓰기가 있으면 따옴표로 감싸세요
  python naver_news_crawler.py                       # 아무것도 안 주면 기본 URL('반도체' 통합검색)
  python naver_news_crawler.py --limit 5 --out news  # 5개만 수집해 news.json / news.csv 로 저장
  python naver_news_crawler.py --limit 5 --xlsx      # 엑셀(naver_news.xlsx)까지 저장
  python naver_news_crawler.py "<네이버 검색 결과 URL>"   # URL을 직접 줘도 됩니다

주의 (꼭 읽어 주세요)
  - 네이버 robots.txt는 일반 봇의 수집을 허용하지 않고(Disallow: /), 서비스 약관도 자동 수집을 제한합니다.
    학습·개인 연구용으로 소량만 사용하고, 수집한 기사를 재배포하거나 AI 학습에 쓰지 마세요.
    기사 저작권은 각 언론사에 있습니다.
  - 그래서 기본값은 소량(10개) + 요청 사이 대기(1초 이상)입니다. 대량 수집에는 쓰지 마세요.
  - 대량/정기 수집이 필요하면 공식 '네이버 검색 API(뉴스)'를 사용하는 것이 안전합니다.
  - 네이버가 화면 구조를 바꾸면 셀렉터를 수정해야 할 수 있습니다.
"""
import argparse
import csv
import json
import random
import re
import sys
import time
from datetime import datetime
from urllib.parse import quote

import requests
from bs4 import BeautifulSoup

DEFAULT_URL = (
    "https://search.naver.com/search.naver?where=nexearch&sm=top_hty&fbm=0&ie=utf8"
    "&query=%EB%B0%98%EB%8F%84%EC%B2%B4&ackey=6zjfxkkq"
)
# 검색어만 입력했을 때 사용할 주소: 네이버 '뉴스' 탭 검색 (통합검색보다 기사 링크가 훨씬 많다)
SEARCH_URL_TEMPLATE = "https://search.naver.com/search.naver?where=news&query={}"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "ko-KR,ko;q=0.9,en;q=0.6",
}
# https://n.news.naver.com/mnews/article/003/0014203493?sid=102  (언론사 코드 / 기사 번호)
ARTICLE_RE = re.compile(r"^https?://n\.news\.naver\.com/(?:mnews/)?article/(\d+)/(\d+)")


def build_search_url(text):
    """검색어 또는 URL을 받아 크롤링할 검색 URL을 돌려준다. 비어 있으면 빈 문자열.

    - 'http(s)://' 로 시작하면 URL로 보고 그대로 사용한다.
    - 그 밖에는 검색어로 보고 네이버 뉴스 검색 URL을 만든다. (한글·공백·특수문자는 자동으로 인코딩)
    """
    text = (text or "").strip()
    if not text:
        return ""
    if re.match(r"^https?://", text, re.IGNORECASE):
        return text
    return SEARCH_URL_TEMPLATE.format(quote(text, safe=""))


def fetch(session, url, retries=3, timeout=10, log=None):
    """URL의 HTML을 가져온다. 실패하면 대기 시간을 늘려 재시도하고, 끝내 실패하면 None.

    log: 실패 메시지를 받을 함수 (GUI 앱에서 사용). 없으면 표준 에러로 출력한다.
    """
    log = log or (lambda msg: print(f"  ! {msg}", file=sys.stderr))
    for attempt in range(1, retries + 1):
        try:
            res = session.get(url, timeout=timeout)
            res.raise_for_status()
            return res.text
        except requests.RequestException as e:
            log(f"요청 실패({attempt}/{retries}): {e}")
            if attempt < retries:
                time.sleep(2 * attempt)
    return None


def extract_article_links(html):
    """검색 결과 HTML에서 네이버 뉴스 기사 URL을 (중복 없이, 등장 순서대로) 뽑는다."""
    soup = BeautifulSoup(html, "html.parser")
    links, seen = [], set()
    for a in soup.find_all("a", href=True):
        m = ARTICLE_RE.match(a["href"])
        if not m or m.groups() in seen:
            continue
        seen.add(m.groups())
        links.append(f"https://n.news.naver.com/mnews/article/{m.group(1)}/{m.group(2)}")
    return links


def clean_body(tag):
    """본문 태그에서 사진 설명·스크립트 등을 걷어내고 문단 단위 텍스트로 만든다."""
    for junk in tag.select("script, style, em.img_desc, span.end_photo_org, .vod_area, .media_end_summary"):
        junk.decompose()
    for br in tag.find_all("br"):
        br.replace_with("\n")
    lines = (line.strip() for line in tag.get_text().splitlines())
    return "\n".join(line for line in lines if line)


def parse_article(html, url):
    """기사 페이지 HTML에서 필요한 정보를 추출한다. 본문을 찾지 못하면 None."""
    soup = BeautifulSoup(html, "html.parser")

    body_tag = soup.select_one("#dic_area") or soup.select_one("#newsct_article")
    if body_tag is None:      # 스포츠·연예 등 다른 형식의 기사는 건너뛴다
        return None

    def text_of(selector):
        el = soup.select_one(selector)
        return el.get_text(" ", strip=True) if el else ""

    title = text_of("#title_area") or text_of("h2.media_end_head_headline")
    if not title:
        og = soup.find("meta", property="og:title")
        title = og["content"].strip() if og and og.get("content") else ""

    logo = soup.select_one(".media_end_head_top_logo img")
    press = logo.get("alt", "").strip() if logo else ""
    if not press:
        og_author = soup.find("meta", property="og:article:author")
        press = og_author["content"].split("|")[0].strip() if og_author and og_author.get("content") else ""

    date_tag = soup.select_one("span.media_end_head_info_datestamp_time")
    published = ""
    if date_tag:
        published = date_tag.get("data-date-time") or date_tag.get_text(strip=True)

    return {
        "url": url,
        "title": title,
        "press": press,
        "reporter": text_of(".byline_s"),
        "published": published,
        "body": clean_body(body_tag),
        "crawled_at": datetime.now().isoformat(timespec="seconds"),
    }


def save(articles, out):
    with open(f"{out}.json", "w", encoding="utf-8") as f:
        json.dump(articles, f, ensure_ascii=False, indent=2)
    with open(f"{out}.csv", "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(articles[0].keys()))
        writer.writeheader()
        writer.writerows(articles)


XLSX_COLUMNS = [            # (헤더, 열 너비)
    ("번호", 6), ("언론사", 14), ("제목", 50), ("기자", 24), ("작성시각", 19), ("본문", 90), ("링크", 46), ("수집시각", 19),
]
EXCEL_CELL_LIMIT = 32000    # 엑셀 셀 하나에 넣을 수 있는 글자 수(32,767)보다 조금 작게


def _to_datetime(text):
    try:
        return datetime.strptime(text, "%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError):
        try:
            return datetime.fromisoformat(text)
        except (TypeError, ValueError):
            return None


def save_xlsx(articles, path):
    """기사 목록을 서식이 적용된 엑셀(.xlsx) 파일로 저장한다. (openpyxl 필요: pip install openpyxl)

    - 헤더 고정 + 자동 필터, 열 너비 지정, 본문 줄바꿈
    - 제목 셀에 기사 링크(하이퍼링크) 연결
    - 작성시각/수집시각은 엑셀 날짜 형식으로 저장해 정렬·필터가 가능하다.
    """
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    def clean(value):
        text = ILLEGAL_CHARACTERS_RE.sub("", value or "")      # 엑셀이 저장을 거부하는 제어 문자 제거
        if len(text) > EXCEL_CELL_LIMIT:
            text = text[:EXCEL_CELL_LIMIT] + "\n…(길이 제한으로 이하 생략)"
        return text

    wb = Workbook()
    ws = wb.active
    ws.title = "뉴스"

    head_font = Font(name="맑은 고딕", bold=True, color="FAF7F0")
    head_fill = PatternFill("solid", fgColor="16130F")
    line = Side(style="thin", color="CFC7B7")
    border = Border(bottom=line)

    for col, (name, width) in enumerate(XLSX_COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 26

    body_font = Font(name="맑은 고딕", size=10)
    top = Alignment(vertical="top")
    wrap = Alignment(vertical="top", wrap_text=True)

    for row, a in enumerate(articles, 2):
        values = [
            row - 1, clean(a.get("press")), clean(a.get("title")), clean(a.get("reporter")),
            _to_datetime(a.get("published")) or clean(a.get("published")),
            clean(a.get("body")), a.get("url", ""),
            _to_datetime(a.get("crawled_at")) or clean(a.get("crawled_at")),
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font, cell.border = body_font, border
            cell.alignment = wrap if col in (3, 6) else top
            if col in (5, 8) and not isinstance(value, str):
                cell.number_format = "yyyy-mm-dd hh:mm"
        if a.get("url"):                                           # 제목 클릭 시 기사 열기
            title_cell = ws.cell(row=row, column=3)
            title_cell.hyperlink = a["url"]
            title_cell.font = Font(name="맑은 고딕", size=10, bold=True, color="A4262C", underline="single")
        ws.row_dimensions[row].height = 105                       # 본문은 셀을 눌러 전체를 볼 수 있다

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(XLSX_COLUMNS))}{len(articles) + 1}"
    wb.save(path)


def main():
    if hasattr(sys.stdout, "reconfigure"):      # 윈도우 콘솔에서 한글 출력이 깨지지 않게
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")

    parser = argparse.ArgumentParser(description="네이버 검색 결과의 뉴스 기사 본문 크롤러")
    parser.add_argument("query", nargs="?", default=DEFAULT_URL,
                        help="검색어 또는 네이버 검색 결과 URL (기본: '반도체' 통합검색 URL)")
    parser.add_argument("--limit", type=int, default=10, help="수집할 최대 기사 수 (기본 10)")
    parser.add_argument("--delay", type=float, default=1.0, help="기사 요청 사이 최소 대기 시간(초). 기본 1.0")
    parser.add_argument("--out", default="naver_news", help="저장 파일 이름(확장자 제외). 기본 naver_news")
    parser.add_argument("--xlsx", action="store_true", help="엑셀(.xlsx) 파일로도 저장 (openpyxl 필요)")
    args = parser.parse_args()

    session = requests.Session()
    session.headers.update(HEADERS)

    url = build_search_url(args.query)
    print(f"검색 주소: {url}")
    print("검색 결과 페이지를 가져오는 중...")
    html = fetch(session, url)
    if html is None:
        sys.exit("검색 결과 페이지를 가져오지 못했습니다.")

    links = extract_article_links(html)
    if not links:
        sys.exit("네이버 뉴스 기사 링크를 찾지 못했습니다. (다른 검색어를 시도하거나, 화면 구조가 바뀌지 않았는지 확인해 주세요)")
    links = links[: args.limit]
    print(f"기사 링크 {len(links)}개를 찾았습니다.\n")

    articles = []
    for i, link in enumerate(links, 1):
        if i > 1:
            time.sleep(args.delay + random.uniform(0, 0.7))      # 서버에 부담을 주지 않도록 대기
        print(f"[{i}/{len(links)}] {link}")
        page = fetch(session, link)
        article = parse_article(page, link) if page else None
        if article is None:
            print("  - 건너뜀 (가져오기 실패 또는 지원하지 않는 기사 형식)")
            continue
        print(f"  - {article['press']} | {article['published']} | {article['title']}")
        articles.append(article)

    if not articles:
        sys.exit("수집된 기사가 없습니다.")

    save(articles, args.out)
    saved = f"{args.out}.json / {args.out}.csv"
    if args.xlsx:
        try:
            save_xlsx(articles, f"{args.out}.xlsx")
            saved += f" / {args.out}.xlsx"
        except ImportError:
            print("엑셀 저장에는 openpyxl 이 필요합니다: pip install openpyxl", file=sys.stderr)
        except OSError as e:
            print(f"엑셀 저장 실패(파일이 엑셀에서 열려 있지 않은지 확인하세요): {e}", file=sys.stderr)
    print(f"\n완료: {len(articles)}개 기사를 {saved} 로 저장했습니다.")

    first = articles[0]
    print("\n--- 첫 번째 기사 미리보기 ---")
    print(first["title"])
    print(first["body"][:300] + ("..." if len(first["body"]) > 300 else ""))


if __name__ == "__main__":
    main()
