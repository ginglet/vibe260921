"""코스피200 '편입종목 상위' 크롤러 (네이버 증권 / Npay 증권)

  https://stock.naver.com/domestic/index/KPI200/price   (코스피200 지수 페이지)

왜 BeautifulSoup4 를 쓰지 않나요?
  요청하신 https://stock.naver.com/market/stock/kr 는 브라우저에서 자바스크립트가 데이터를 채우는 페이지라,
  받아 온 HTML 안에 종목 표(편입종목)가 아예 없습니다. 예전에 HTML 표로 제공하던
  finance.naver.com/sise/entryJongmok.naver 도 '더 이상 제공되지 않음(410)'으로 바뀌었습니다.
  그래서 이 페이지가 화면을 그릴 때 내부적으로 호출하는 JSON 주소(enrollStocks)를 requests 로 직접 읽습니다.
  (JSON 은 HTML 이 아니라서 BeautifulSoup 이 할 일이 없습니다.)

사용법
  pip install requests openpyxl                      # openpyxl 은 엑셀 저장(--xlsx)을 쓸 때만 필요
  python kospi200_crawler.py                         # 코스피200 전 종목(약 200개) 수집, 상위 10개 출력
  python kospi200_crawler.py --limit 50 --xlsx       # 상위 50개만, 엑셀까지 저장
  python kospi200_crawler.py --out kospi200_today    # kospi200_today.json / .csv 로 저장

주의 (꼭 읽어 주세요)
  - stock.naver.com 의 robots.txt 는 일반 봇에 대해 Disallow: / 이고, 이 JSON 주소는 문서화된 공식 API 가 아닙니다.
    개인 학습·연구용으로 소량만 쓰세요. (전체 수집 = 페이지 요청 약 5회 + 지수 요청 1회, 요청 사이 대기 포함)
    주소·응답 형식은 예고 없이 바뀌거나 차단될 수 있습니다.
  - 안정적으로 쓰려면 한국거래소(KRX) 정보데이터시스템 또는 공공데이터포털(금융위원회 지수시세정보) 같은
    공식 데이터를 사용하는 것이 좋습니다. 시세 데이터의 저작권은 각 제공처에 있습니다.
  - 투자 판단의 근거로 쓰지 마세요. 데이터는 지연되거나 틀릴 수 있습니다.
"""
import argparse
import csv
import json
import sys
import time
from datetime import datetime

import requests

INDEX_CODE = "KPI200"                 # 코스피200
API_BASE = "https://stock.naver.com/api/securityFe/api/index"
PAGE_SIZE = 50                        # 서버가 허용하는 최대는 60 (70 이상은 400 오류)
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept": "application/json, text/plain, */*",
    "Accept-Language": "ko-KR,ko;q=0.9,en;q=0.6",
    "Referer": f"https://stock.naver.com/domestic/index/{INDEX_CODE}/price",
}


class CrawlError(Exception):
    """수집 실패(네트워크 오류, 응답 형식 변경 등)."""


# ---------------------------------------------------------------- 파싱 도우미
def to_int(text):
    """'1,872,000' / '+12,000' / '-10,000' → int. 값이 없거나 'N/A' 이면 None."""
    try:
        return int(str(text).replace(",", "").replace("+", "").strip())
    except (TypeError, ValueError):
        return None


def to_float(text):
    try:
        return float(str(text).replace(",", "").replace("+", "").replace("%", "").strip())
    except (TypeError, ValueError):
        return None


def direction_of(raw):
    """등락 방향을 'up' / 'down' / 'flat' 으로 통일한다. (RISING, UPPER_LIMIT, FALLING, LOWER_LIMIT, UNCHANGED …)"""
    name = ((raw or {}).get("name") or "").upper()
    if "RIS" in name or "UPPER" in name:
        return "up"
    if "FALL" in name or "LOWER" in name:
        return "down"
    return "flat"


def normalize(raw, rank):
    """API 응답 한 종목을 화면·파일에 쓰기 좋은 평평한 dict 로 바꾼다.

    단위: 거래량=주, 거래대금=백만원, 시가총액=억원
    """
    direction = direction_of(raw.get("compareToPreviousPrice"))
    sign = {"up": 1, "down": -1, "flat": 0}[direction]
    change = to_int(raw.get("compareToPreviousClosePrice"))
    rate = to_float(raw.get("fluctuationsRatio"))
    return {
        "rank": rank,
        "code": raw.get("itemCode", ""),
        "name": raw.get("stockName", ""),
        "price": to_int(raw.get("closePrice")),
        "change": abs(change) * sign if change is not None else None,
        "change_rate": abs(rate) * sign if rate is not None else None,       # 단위: %
        "direction": direction,
        "volume": to_int(raw.get("accumulatedTradingVolume")),
        "trade_value": to_int(raw.get("accumulatedTradingValue")),
        "market_cap": to_int(raw.get("marketValue")),
        "market": (raw.get("stockExchangeType") or {}).get("nameKor", ""),
        "traded_at": raw.get("localTradedAt", ""),
        "url": raw.get("newPcUrl") or f"https://stock.naver.com/domestic/stock/{raw.get('itemCode', '')}",
    }


# ---------------------------------------------------------------- 요청
def get_json(session, url, retries=3, timeout=10, log=None):
    """JSON 을 가져온다. 실패하면 대기 시간을 늘려 재시도하고, 끝내 실패하면 CrawlError."""
    log = log or (lambda msg: print(f"  ! {msg}", file=sys.stderr))
    last = None
    for attempt in range(1, retries + 1):
        try:
            res = session.get(url, timeout=timeout)
            res.raise_for_status()
            return res.json()
        except (requests.RequestException, ValueError) as e:
            last = e
            log(f"요청 실패({attempt}/{retries}): {e}")
            if attempt < retries:
                time.sleep(2 * attempt)
    raise CrawlError(f"데이터를 가져오지 못했습니다: {last}")


def new_session():
    session = requests.Session()
    session.headers.update(HEADERS)
    return session


def fetch_index_summary(session, index_code=INDEX_CODE, log=None):
    """지수의 최근 종가·전일비·등락률. 실패해도 종목 수집은 계속할 수 있도록 None 을 돌려준다."""
    try:
        data = get_json(session, f"{API_BASE}/{index_code}/price?page=1&pageSize=1", retries=2, log=log)
        row = data[0]
        direction = direction_of(row.get("compareToPreviousPrice"))
        sign = {"up": 1, "down": -1, "flat": 0}[direction]
        change = to_float(row.get("compareToPreviousClosePrice"))
        rate = to_float(row.get("fluctuationsRatio"))
        return {
            "level": to_float(row.get("closePrice")),
            "change": abs(change) * sign if change is not None else None,
            "change_rate": abs(rate) * sign if rate is not None else None,
            "direction": direction,
            "date": row.get("localTradedAt", ""),
        }
    except (CrawlError, IndexError, KeyError, TypeError):
        return None


def fetch_constituents(limit=None, delay=1.0, index_code=INDEX_CODE, session=None,
                       log=None, should_stop=None, on_progress=None, on_page=None):
    """편입종목 상위 목록을 순서대로(순위 1부터) 가져온다.

    limit        : 가져올 최대 종목 수 (None 이면 전부, 코스피200 은 약 200개)
    delay        : 페이지 요청 사이 대기(초)
    should_stop  : True 를 돌려주면 중단하고 지금까지 모은 것을 돌려준다
    on_progress  : on_progress(지금까지 모은 개수) 콜백
    on_page      : on_page(이번 페이지에서 새로 모은 종목 리스트) 콜백 - 화면에 바로바로 보여 줄 때 사용
    """
    session = session or new_session()
    rows, page = [], 1
    while limit is None or len(rows) < limit:
        if should_stop and should_stop():
            break
        if page > 1:
            end = time.monotonic() + delay
            while time.monotonic() < end and not (should_stop and should_stop()):
                time.sleep(0.05)                       # 대기 중에도 '중지'에 바로 반응하도록 잘게 쉰다
            if should_stop and should_stop():
                break
        url =f"{API_BASE}/{index_code}/enrollStocks?page={page}&pageSize={PAGE_SIZE}&type=list"
        data = get_json(session, url, log=log)
        if not isinstance(data, list):
            raise CrawlError("응답 형식이 예상과 다릅니다. (사이트 구조가 바뀌었을 수 있습니다)")
        start = len(rows)
        for raw in data:
            if limit is not None and len(rows) >= limit:
                break
            rows.append(normalize(raw, len(rows) + 1))
        if on_page and len(rows) > start:
            on_page(rows[start:])
        if on_progress:
            on_progress(len(rows))
        if len(data) < PAGE_SIZE:          # 마지막 페이지
            break
        page += 1
    return rows


# ---------------------------------------------------------------- 저장
FIELDS = ["rank", "code", "name", "price", "change", "change_rate", "direction",
          "volume", "trade_value", "market_cap", "market", "traded_at", "url"]


def save_json(rows, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)


def save_csv(rows, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:      # 엑셀에서 한글이 깨지지 않도록 BOM 포함
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)


XLSX_COLUMNS = [       # (헤더, 열 너비, dict 키, 엑셀 숫자 서식)
    ("순위", 7, "rank", "0"),
    ("종목명", 22, "name", None),
    ("종목코드", 10, "code", None),
    ("현재가(원)", 13, "price", "#,##0"),
    ("전일비(원)", 12, "change", "+#,##0;-#,##0;0"),
    ("등락률", 10, "change_rate", "+0.00%;-0.00%;0.00%"),
    ("거래량(주)", 15, "volume", "#,##0"),
    ("거래대금(백만원)", 17, "trade_value", "#,##0"),
    ("시가총액(억원)", 16, "market_cap", "#,##0"),
    ("시장", 8, "market", None),
    ("기준시각", 19, "traded_at", "yyyy-mm-dd hh:mm"),
]
UP_COLOR, DOWN_COLOR = "D32F2F", "1565C0"      # 한국 증시 관례: 상승 빨강 / 하락 파랑


def save_xlsx(rows, path, index_summary=None):
    """서식이 적용된 엑셀(.xlsx)로 저장한다. (openpyxl 필요: pip install openpyxl)

    - 숫자는 숫자로 저장 → 엑셀에서 정렬·필터·수식 사용 가능
    - 상승은 빨강, 하락은 파랑, 등락률은 % 서식
    - 종목명은 종목 페이지 하이퍼링크
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "코스피200"

    head_font = Font(name="맑은 고딕", bold=True, color="FAF7F0")
    head_fill = PatternFill("solid", fgColor="16130F")
    border = Border(bottom=Side(style="thin", color="CFC7B7"))
    base = Font(name="맑은 고딕", size=10)

    for col, (name, width, _, _) in enumerate(XLSX_COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 26

    for r, row in enumerate(rows, 2):
        color = {"up": UP_COLOR, "down": DOWN_COLOR}.get(row["direction"])
        for col, (_, _, key, fmt) in enumerate(XLSX_COLUMNS, 1):
            value = row.get(key)
            if key == "change_rate" and value is not None:
                value = value / 100                                  # 엑셀 % 서식용
            if key == "traded_at":
                try:
                    value = datetime.fromisoformat(value).replace(tzinfo=None)
                except (TypeError, ValueError):
                    pass
            cell = ws.cell(row=r, column=col, value=value)
            cell.font, cell.border = base, border
            if fmt:
                cell.number_format = fmt
            if key in ("price", "change", "change_rate") and color:
                cell.font = Font(name="맑은 고딕", size=10, color=color)
            if key in ("rank", "code", "market"):
                cell.alignment = Alignment(horizontal="center")
        name_cell = ws.cell(row=r, column=2)
        if row.get("url"):
            name_cell.hyperlink = row["url"]
            name_cell.font = Font(name="맑은 고딕", size=10, bold=True, underline="single")

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(XLSX_COLUMNS))}{len(rows) + 1}"

    if index_summary and index_summary.get("level") is not None:       # 지수 요약은 두 번째 시트에
        info = wb.create_sheet("지수 요약")
        pairs = [("지수", "코스피200"), ("종가", index_summary["level"]), ("전일비", index_summary["change"]),
                 ("등락률(%)", index_summary["change_rate"]), ("기준일", index_summary["date"][:10]),
                 ("수집 종목 수", len(rows)), ("수집 시각", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))]
        for i, (k, v) in enumerate(pairs, 1):
            info.cell(row=i, column=1, value=k).font = Font(name="맑은 고딕", bold=True)
            info.cell(row=i, column=2, value=v).font = Font(name="맑은 고딕")
        info.column_dimensions["A"].width = 14
        info.column_dimensions["B"].width = 22
    wb.save(path)


# ---------------------------------------------------------------- 명령줄
def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")

    parser = argparse.ArgumentParser(description="코스피200 편입종목 상위 크롤러")
    parser.add_argument("--limit", type=int, default=None, help="가져올 최대 종목 수 (기본: 전부)")
    parser.add_argument("--delay", type=float, default=1.0, help="페이지 요청 사이 대기 시간(초). 기본 1.0, 최소 0.5")
    parser.add_argument("--out", default="kospi200", help="저장 파일 이름(확장자 제외). 기본 kospi200")
    parser.add_argument("--xlsx", action="store_true", help="엑셀(.xlsx) 파일로도 저장 (openpyxl 필요)")
    args = parser.parse_args()
    args.delay = max(0.5, args.delay)

    session = new_session()
    print("코스피200 지수 정보를 가져오는 중...")
    summary = fetch_index_summary(session)
    if summary:
        arrow = {"up": "▲", "down": "▼", "flat": "-"}[summary["direction"]]
        print(f"  코스피200 {summary['level']:,.2f}  {arrow}{abs(summary['change'] or 0):,.2f} ({summary['change_rate']:+.2f}%)  기준일 {summary['date'][:10]}")

    print("편입종목 상위 목록을 가져오는 중...")
    try:
        rows = fetch_constituents(limit=args.limit, delay=args.delay, session=session,
                                  on_progress=lambda n: print(f"  {n}개 수집"))
    except CrawlError as e:
        sys.exit(str(e))
    if not rows:
        sys.exit("수집된 종목이 없습니다.")

    save_json(rows, f"{args.out}.json")
    save_csv(rows, f"{args.out}.csv")
    saved = f"{args.out}.json / {args.out}.csv"
    if args.xlsx:
        try:
            save_xlsx(rows, f"{args.out}.xlsx", summary)
            saved += f" / {args.out}.xlsx"
        except ImportError:
            print("엑셀 저장에는 openpyxl 이 필요합니다: pip install openpyxl", file=sys.stderr)
        except OSError as e:
            print(f"엑셀 저장 실패(파일이 엑셀에서 열려 있지 않은지 확인하세요): {e}", file=sys.stderr)
    print(f"\n완료: {len(rows)}개 종목을 {saved} 로 저장했습니다.\n")

    print(f"{'순위':>4}  {'종목명':<14}{'현재가':>11}{'전일비':>10}{'등락률':>8}{'시가총액(억)':>14}")
    for r in rows[:10]:
        change = f"{r['change']:+,}" if r["change"] is not None else "-"
        rate = f"{r['change_rate']:+.2f}%" if r["change_rate"] is not None else "-"
        price = f"{r['price']:,}" if r["price"] is not None else "-"
        cap = f"{r['market_cap']:,}" if r["market_cap"] is not None else "-"
        print(f"{r['rank']:>4}  {r['name']:<14}{price:>11}{change:>10}{rate:>8}{cap:>14}")


if __name__ == "__main__":
    main()
